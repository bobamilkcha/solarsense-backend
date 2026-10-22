"""Export a FeasibilityReport as an Excel workbook (openpyxl).

Three sheets: Summary (headline figures), Monthly (month-by-month table with
SUM formulas for the year), Assumptions (inputs used). Values are written as
numbers with number formats, not pre-formatted strings, so they stay usable
for the installer's own calculations.
"""

import io
from datetime import datetime

from app.modules.reports.feasibility import Assumptions
from app.modules.reports.schemas import FeasibilityReport
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

RM = '"RM" #,##0.00'
RM0 = '"RM" #,##0'
KWH = "#,##0"
PCT = "0%"
DEC1 = "0.0"

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
MUTED = Font(color="616E7C")
HEADER_FILL = PatternFill("solid", fgColor="E8EDF2")
TOP_RULE = Border(top=Side(style="thin"))


def _rows(ws: Worksheet, start_row: int, rows: list[tuple[str, object, str | None]]) -> int:
    """Write label/value rows; returns the next free row."""
    for i, (label, value, fmt) in enumerate(rows):
        row = start_row + i
        ws.cell(row=row, column=1, value=label).font = MUTED
        cell = ws.cell(row=row, column=2, value=value)
        if fmt:
            cell.number_format = fmt
    return start_row + len(rows)


def _summary(ws: Worksheet, r: FeasibilityReport, generated_at: datetime) -> None:
    ws.title = "Summary"
    ws["A1"] = "Solar Feasibility Report"
    ws["A1"].font = TITLE
    ws["A2"] = f"{r.site_name} · {r.location_lat:.4f}, {r.location_lng:.4f}"
    ws["A3"] = f"Generated {generated_at:%d %b %Y %H:%M} UTC"
    ws["A3"].font = MUTED

    c, s = r.consumption, r.shading
    row = _rows(
        ws,
        5,
        [
            ("Recommended system (kWp)", r.system_kwp, "0.00"),
            ("Panels", r.panel_count, None),
            ("Estimated system cost", r.system_cost_rm, RM0),
            ("Annual solar output (kWh)", r.annual_production_kwh, KWH),
            ("Annual usage (kWh)", r.annual_consumption_kwh, KWH),
            ("Usage covered by solar", r.solar_coverage_pct / 100, PCT),
            ("Average monthly bill without solar", r.avg_monthly_bill_before_rm, RM),
            ("Average monthly bill with solar", r.avg_monthly_bill_after_rm, RM),
            ("Annual savings", r.annual_savings_rm, RM0),
            ("Simple payback (years)", r.payback_years, DEC1),
            (f"Net savings over {r.lifetime_years} years", r.lifetime_net_savings_rm, RM0),
        ],
    )
    ws.cell(row=row + 1, column=1, value="Site").font = BOLD
    _rows(
        ws,
        row + 2,
        [
            ("Bills used", c.bills_used, None),
            ("Average monthly usage (kWh)", c.avg_monthly_kwh, KWH),
            ("Tariff", c.tariff_category, None),
            ("Bill calculation", c.tariff_model, None),
            ("Effective rate from bills (RM/kWh)", c.effective_rate_rm_per_kwh, "0.0000"),
            ("Satellite sunlight (kWh/m² per year)", r.annual_irradiation_kwh_m2, KWH),
            ("Shading source", s.source, None),
            ("Full days measured on-site", s.complete_days, None),
            ("Measured vs satellite", s.measured_ratio, PCT),
            ("Shading factor applied", s.applied_factor, PCT),
        ],
    )
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 18


def _monthly(ws: Worksheet, r: FeasibilityReport) -> None:
    headers = [
        ("Month", None),
        ("Solar output (kWh)", KWH),
        ("Used on-site (kWh)", KWH),
        ("Exported (kWh)", KWH),
        ("Bill without solar", RM),
        ("Bill with solar", RM),
        ("Savings", RM),
    ]
    for col, (name, _) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = BOLD
        cell.fill = HEADER_FILL

    for i, m in enumerate(r.months):
        values = [
            MONTHS[m.month - 1],
            m.production_kwh,
            m.self_consumed_kwh,
            m.exported_kwh,
            m.bill_before_rm,
            m.bill_after_rm,
            m.savings_rm,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=i + 2, column=col, value=value)
            if fmt := headers[col - 1][1]:
                cell.number_format = fmt

    last = len(r.months) + 1
    total_row = last + 1
    ws.cell(row=total_row, column=1, value="Year").font = BOLD
    for col in range(2, len(headers) + 1):
        letter = get_column_letter(col)
        cell = ws.cell(row=total_row, column=col, value=f"=SUM({letter}2:{letter}{last})")
        cell.font = BOLD
        cell.border = TOP_RULE
        cell.number_format = headers[col - 1][1] or "General"
    ws.cell(row=total_row, column=1).border = TOP_RULE

    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 8
    for col in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 19


def _assumptions(ws: Worksheet, r: FeasibilityReport) -> None:
    a = r.assumptions
    ws["A1"] = "Input"
    ws["B1"] = "Value"
    for cell in (ws["A1"], ws["B1"]):
        cell.font = BOLD
        cell.fill = HEADER_FILL
    _rows(
        ws,
        2,
        [
            ("System size (kWp)", a.system_kwp, "0.00"),
            ("Panel size (W)", a.panel_wp, None),
            ("Installed cost (RM per kWp)", a.cost_per_kwp_rm, RM0),
            ("Performance ratio", a.performance_ratio, PCT),
            ("Solar used on-site", a.self_consumption_ratio, PCT),
            ("Export credit (RM/kWh)", a.export_rate_rm_per_kwh, "0.0000"),
            ("AFA fuel adjustment (RM/kWh)", a.afa_rate_rm_per_kwh, "0.0000"),
            ("Panel degradation per year", Assumptions().degradation_per_year, "0.0%"),
            ("Lifetime (years)", r.lifetime_years, None),
        ],
    )
    ws["A12"] = (
        "Estimates for planning, not a quotation. Sunlight: PVGIS long-term monthly "
        "averages (horizontal surface). Usage assumed constant through the year; tariffs "
        "held flat. Excludes TNB's Energy Efficiency Incentive rebate."
    )
    ws["A12"].font = MUTED
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14


def render_feasibility_xlsx(r: FeasibilityReport, generated_at: datetime) -> bytes:
    wb = Workbook()
    summary = wb.active
    assert summary is not None  # a new Workbook always has one sheet
    _summary(summary, r, generated_at)
    _monthly(wb.create_sheet("Monthly"), r)
    _assumptions(wb.create_sheet("Assumptions"), r)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
