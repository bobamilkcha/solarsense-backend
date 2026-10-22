"""Integration tests for the feasibility report (service + HTTP), PVGIS mocked."""

import io
import uuid
from datetime import date, datetime, timedelta
from typing import Any

import pytest
from app.core.config.app_config import app_config
from app.models import TNBBill
from app.modules.benchmark.client import PVGISError
from app.modules.benchmark.irradiation import LOCAL_TZ
from app.modules.readings.schemas import ReadingPayload
from app.modules.readings.service import ReadingsService
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession
from tests.integration.test_bills_routes import _auth_headers, _own_site

FLAT_IRR = dict.fromkeys(range(1, 13), 150.0)


@pytest.fixture
def fake_pvgis(monkeypatch: pytest.MonkeyPatch) -> None:
    async def fake_fetch(lat, lng, *, client=None):
        return FLAT_IRR

    monkeypatch.setattr(
        "app.modules.reports.service.fetch_monthly_irradiation_averages", fake_fetch
    )


async def _add_bill(db: AsyncSession, site_id: uuid.UUID, **fields: Any) -> None:
    defaults: dict[str, Any] = {
        "object_key": "k",
        "filename": "bill.pdf",
        "billing_period_start": date(2026, 4, 4),
        "billing_period_end": date(2026, 5, 3),
        "kwh_used": 1295.0,
        "tariff_category": "Domestik Am",
        "current_charges": 616.22,
    }
    db.add(TNBBill(site_id=site_id, **(defaults | fields)))
    await db.commit()


async def test_feasibility_requires_auth(client: AsyncClient) -> None:
    resp = await client.get(f"/sites/{uuid.uuid4()}/feasibility")
    assert resp.status_code == 401


async def test_feasibility_without_bills_422s(
    client: AsyncClient, db: AsyncSession, registered: dict[str, Any], fake_pvgis: None
) -> None:
    site = await _own_site(db, registered)
    resp = await client.get(f"/sites/{site.id}/feasibility", headers=_auth_headers(registered))
    assert resp.status_code == 422


async def test_feasibility_unknown_site_404s(
    client: AsyncClient, registered: dict[str, Any]
) -> None:
    resp = await client.get(f"/sites/{uuid.uuid4()}/feasibility", headers=_auth_headers(registered))
    assert resp.status_code == 404


async def test_feasibility_domestic_bill_satellite_only(
    client: AsyncClient, db: AsyncSession, registered: dict[str, Any], fake_pvgis: None
) -> None:
    site = await _own_site(db, registered)
    await _add_bill(db, site.id)

    resp = await client.get(
        f"/sites/{site.id}/feasibility",
        params={"system_kwp": 5, "cost_per_kwp_rm": 4000},
        headers=_auth_headers(registered),
    )
    assert resp.status_code == 200
    body = resp.json()

    # 30-day bill normalised to an average 30.44-day month
    assert body["consumption"]["avg_monthly_kwh"] == pytest.approx(1295 * 30.4375 / 30)
    assert body["consumption"]["tariff_model"] == "domestic_2025"
    assert body["shading"]["source"] == "satellite_only"
    assert body["shading"]["applied_factor"] == 1.0
    assert body["system_kwp"] == 5
    assert body["system_cost_rm"] == pytest.approx(20000)
    # 5 kWp * 1800 kWh/m^2 * 0.8
    assert body["annual_production_kwh"] == pytest.approx(7200)
    assert body["annual_savings_rm"] > 0
    assert body["payback_years"] == pytest.approx(20000 / body["annual_savings_rm"])
    assert len(body["months"]) == 12


async def test_feasibility_recommends_size_when_omitted(
    client: AsyncClient, db: AsyncSession, registered: dict[str, Any], fake_pvgis: None
) -> None:
    site = await _own_site(db, registered)
    await _add_bill(db, site.id)

    resp = await client.get(f"/sites/{site.id}/feasibility", headers=_auth_headers(registered))
    body = resp.json()

    assert body["assumptions"]["system_kwp"] == body["system_kwp"]
    # Sized to roughly cover annual usage
    assert body["annual_production_kwh"] == pytest.approx(body["annual_consumption_kwh"], rel=0.1)


async def test_feasibility_non_domestic_uses_effective_rate(
    client: AsyncClient, db: AsyncSession, registered: dict[str, Any], fake_pvgis: None
) -> None:
    site = await _own_site(db, registered)
    await _add_bill(
        db, site.id, tariff_category="Komersial", kwh_used=1000.0, current_charges=500.0
    )

    resp = await client.get(f"/sites/{site.id}/feasibility", headers=_auth_headers(registered))
    body = resp.json()

    assert body["consumption"]["tariff_model"] == "flat_effective_rate"
    assert body["consumption"]["effective_rate_rm_per_kwh"] == pytest.approx(0.5)


async def test_feasibility_applies_device_shading(
    client: AsyncClient,
    db: AsyncSession,
    registered: dict[str, Any],
    fake_pvgis: None,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(app_config, "REFERENCE_CELL_ISC_A", 0.2)
    site = await _own_site(db, registered)
    await _add_bill(db, site.id)

    # One complete day at 200 W/m^2 for 12 h = 2.4 kWh/m^2, vs satellite
    # 150/31 = 4.84/day in March -> ratio ~0.496.
    readings = ReadingsService(db)
    t = datetime(2026, 3, 15, 7, 0, tzinfo=LOCAL_TZ)
    while t <= datetime(2026, 3, 15, 19, 0, tzinfo=LOCAL_TZ):
        await readings.record_reading(
            site.device_id,
            ReadingPayload(
                timestamp=t,
                cell_voltage=0.4,  # 0.04 A -> 200 W/m^2
                ambient_temp=30,
                humidity=70,
                tilt_x=0,
                tilt_y=0,
                tilt_z=90,
                heading=180,
            ),
        )
        t += timedelta(minutes=10)

    resp = await client.get(
        f"/sites/{site.id}/feasibility",
        params={"system_kwp": 5},
        headers=_auth_headers(registered),
    )
    body = resp.json()

    assert body["shading"]["source"] == "device"
    assert body["shading"]["complete_days"] == 1
    assert body["shading"]["applied_factor"] == pytest.approx(2.4 / (150 / 31), rel=1e-3)
    assert body["annual_production_kwh"] == pytest.approx(
        7200 * body["shading"]["applied_factor"], rel=1e-6
    )


async def test_feasibility_pvgis_down_502s(
    client: AsyncClient,
    db: AsyncSession,
    registered: dict[str, Any],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def fail(lat, lng, *, client=None):
        raise PVGISError("down")

    monkeypatch.setattr("app.modules.reports.service.fetch_monthly_irradiation_averages", fail)
    site = await _own_site(db, registered)
    await _add_bill(db, site.id)

    resp = await client.get(f"/sites/{site.id}/feasibility", headers=_auth_headers(registered))
    assert resp.status_code == 502


async def test_feasibility_xlsx_download(
    client: AsyncClient, db: AsyncSession, registered: dict[str, Any], fake_pvgis: None
) -> None:
    site = await _own_site(db, registered)
    await _add_bill(db, site.id)

    resp = await client.get(
        f"/sites/{site.id}/feasibility/xlsx",
        params={"system_kwp": 5},
        headers=_auth_headers(registered),
    )

    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in resp.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Summary", "Monthly", "Assumptions"]
    assert "Bill Test Site" in wb["Summary"]["A2"].value
    assert wb["Summary"]["B5"].value == 5


async def test_feasibility_xlsx_without_bills_422s(
    client: AsyncClient, db: AsyncSession, registered: dict[str, Any], fake_pvgis: None
) -> None:
    site = await _own_site(db, registered)
    resp = await client.get(f"/sites/{site.id}/feasibility/xlsx", headers=_auth_headers(registered))
    assert resp.status_code == 422
