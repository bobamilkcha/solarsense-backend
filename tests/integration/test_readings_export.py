"""Integration tests for GET /sites/{site_id}/readings/export (CSV + Excel)."""

import csv
import io
import uuid
from datetime import UTC, datetime
from typing import Any

import pytest
from app.core.config.app_config import app_config
from app.modules.readings.schemas import ReadingPayload
from app.modules.readings.service import ReadingsService
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession
from tests.integration.test_bills_routes import _auth_headers, _own_site
from tests.integration.test_readings_service import _make_site

TIMES = [
    datetime(2026, 3, 15, 2, 0, tzinfo=UTC),  # 10:00 in Malaysia
    datetime(2026, 3, 15, 4, 0, tzinfo=UTC),
    datetime(2026, 3, 16, 4, 0, tzinfo=UTC),
]


async def _seed(db: AsyncSession, device_id: str) -> None:
    svc = ReadingsService(db)
    # Insert out of order to check the export sorts oldest first.
    for ts, volts in [(TIMES[2], 1.0), (TIMES[0], 0.4), (TIMES[1], 2.0)]:
        await svc.record_reading(
            device_id,
            ReadingPayload(
                timestamp=ts,
                cell_voltage=volts,
                ambient_temp=31.5,
                humidity=68.2,
                tilt_x=1.0,
                tilt_y=-2.0,
                tilt_z=89.0,
                heading=180.0,
            ),
        )


def _csv_rows(content: bytes) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))


async def test_export_requires_auth(client: AsyncClient) -> None:
    resp = await client.get(f"/sites/{uuid.uuid4()}/readings/export")
    assert resp.status_code == 401


async def test_export_other_company_site_404s(
    client: AsyncClient, db: AsyncSession, registered: dict[str, Any]
) -> None:
    other_site = await _make_site(db)
    resp = await client.get(
        f"/sites/{other_site.id}/readings/export", headers=_auth_headers(registered)
    )
    assert resp.status_code == 404


async def test_export_csv(
    client: AsyncClient,
    db: AsyncSession,
    registered: dict[str, Any],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(app_config, "REFERENCE_CELL_ISC_A", 0.2)
    site = await _own_site(db, registered)
    await _seed(db, site.device_id)

    resp = await client.get(f"/sites/{site.id}/readings/export", headers=_auth_headers(registered))

    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert f"readings-{site.device_id}-" in resp.headers["content-disposition"]
    assert resp.content.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM for Excel

    rows = _csv_rows(resp.content)
    assert len(rows) == 3
    first = rows[0]
    assert datetime.fromisoformat(first["timestamp_utc"]) == TIMES[0]
    assert first["timestamp_local"] == "2026-03-15T10:00:00+08:00"
    assert float(first["cell_voltage_v"]) == pytest.approx(0.4)
    assert float(first["cell_current_a"]) == pytest.approx(0.04)
    assert float(first["irradiance_wm2"]) == pytest.approx(200.0)
    assert [datetime.fromisoformat(r["timestamp_utc"]) for r in rows] == TIMES


async def test_export_date_range_filter(
    client: AsyncClient, db: AsyncSession, registered: dict[str, Any]
) -> None:
    site = await _own_site(db, registered)
    await _seed(db, site.device_id)

    resp = await client.get(
        f"/sites/{site.id}/readings/export",
        params={"start": TIMES[1].isoformat(), "end": TIMES[2].isoformat()},
        headers=_auth_headers(registered),
    )

    rows = _csv_rows(resp.content)
    # start inclusive, end exclusive
    assert [datetime.fromisoformat(r["timestamp_utc"]) for r in rows] == [TIMES[1]]


async def test_export_invalid_range_422s(
    client: AsyncClient, db: AsyncSession, registered: dict[str, Any]
) -> None:
    site = await _own_site(db, registered)
    resp = await client.get(
        f"/sites/{site.id}/readings/export",
        params={"start": TIMES[2].isoformat(), "end": TIMES[0].isoformat()},
        headers=_auth_headers(registered),
    )
    assert resp.status_code == 422


async def test_export_empty_site_gives_header_only(
    client: AsyncClient, db: AsyncSession, registered: dict[str, Any]
) -> None:
    site = await _own_site(db, registered)
    resp = await client.get(f"/sites/{site.id}/readings/export", headers=_auth_headers(registered))
    assert resp.status_code == 200
    assert _csv_rows(resp.content) == []


async def test_export_xlsx(
    client: AsyncClient, db: AsyncSession, registered: dict[str, Any]
) -> None:
    site = await _own_site(db, registered)
    await _seed(db, site.device_id)

    resp = await client.get(
        f"/sites/{site.id}/readings/export",
        params={"format": "xlsx"},
        headers=_auth_headers(registered),
    )

    assert resp.status_code == 200
    assert resp.headers["content-disposition"].endswith('.xlsx"')
    ws = load_workbook(io.BytesIO(resp.content))["Readings"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:3] == ("timestamp_utc", "timestamp_local", "cell_voltage_v")
    assert len(rows) == 4
    # Real Excel datetimes (wall-clock), not strings
    assert rows[1][0] == datetime(2026, 3, 15, 2, 0)
    assert rows[1][1] == datetime(2026, 3, 15, 10, 0)


async def test_export_xlsx_too_large_422s(
    client: AsyncClient,
    db: AsyncSession,
    registered: dict[str, Any],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr("app.modules.readings.router.XLSX_MAX_ROWS", 2)
    site = await _own_site(db, registered)
    await _seed(db, site.device_id)

    resp = await client.get(
        f"/sites/{site.id}/readings/export",
        params={"format": "xlsx"},
        headers=_auth_headers(registered),
    )
    assert resp.status_code == 422
