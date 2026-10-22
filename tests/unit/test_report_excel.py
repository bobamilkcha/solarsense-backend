"""Excel export of the feasibility report."""

import io
import uuid
from datetime import UTC, datetime
from typing import Any

from app.modules.reports.excel import render_feasibility_xlsx
from app.modules.reports.schemas import (
    AssumptionsParams,
    ConsumptionSummary,
    FeasibilityReport,
    MonthBreakdown,
    ShadingSummary,
)
from openpyxl import load_workbook


def _report(**overrides: Any) -> FeasibilityReport:
    fields: dict[str, Any] = {
        "site_id": uuid.uuid4(),
        "site_name": "Test Rooftop",
        "location_lat": 3.139,
        "location_lng": 101.6869,
        "consumption": ConsumptionSummary(
            bills_used=1,
            avg_monthly_kwh=1000,
            tariff_category="Domestik Am",
            tariff_model="domestic_2025",
            effective_rate_rm_per_kwh=0.48,
        ),
        "shading": ShadingSummary(
            source="satellite_only", complete_days=0, measured_ratio=None, applied_factor=1.0
        ),
        "system_kwp": 5.5,
        "panel_count": 10,
        "system_cost_rm": 22000,
        "annual_irradiation_kwh_m2": 1750,
        "annual_production_kwh": 7680,
        "annual_consumption_kwh": 12000,
        "solar_coverage_pct": 38.4,
        "avg_monthly_bill_before_rm": 480,
        "avg_monthly_bill_after_rm": 300,
        "annual_savings_rm": 2160,
        "payback_years": 10.2,
        "lifetime_net_savings_rm": 28000,
        "lifetime_years": 25,
        "assumptions": AssumptionsParams(system_kwp=5.5),
        "months": [
            MonthBreakdown(
                month=m,
                production_kwh=640,
                self_consumed_kwh=384,
                exported_kwh=256,
                bill_before_rm=480,
                bill_after_rm=300,
                savings_rm=180,
            )
            for m in range(1, 13)
        ],
    }
    return FeasibilityReport(**(fields | overrides))


def _load(report: FeasibilityReport):
    data = render_feasibility_xlsx(report, datetime(2026, 10, 22, tzinfo=UTC))
    return load_workbook(io.BytesIO(data))


def test_sheets_and_summary_values() -> None:
    wb = _load(_report())

    assert wb.sheetnames == ["Summary", "Monthly", "Assumptions"]
    summary = wb["Summary"]
    assert "Test Rooftop" in summary["A2"].value
    assert summary["B5"].value == 5.5  # kWp, stored as a number
    assert summary["B7"].value == 22000
    assert summary["B10"].value == 0.384  # coverage stored as a fraction
    assert summary["B10"].number_format == "0%"


def test_monthly_sheet_has_12_months_and_sum_formulas() -> None:
    monthly = _load(_report())["Monthly"]

    assert monthly["A2"].value == "Jan"
    assert monthly["A13"].value == "Dec"
    assert monthly["B2"].value == 640
    assert monthly["A14"].value == "Year"
    assert monthly["B14"].value == "=SUM(B2:B13)"
    assert monthly["G14"].value == "=SUM(G2:G13)"


def test_missing_values_left_blank() -> None:
    summary = _load(_report(payback_years=None))["Summary"]

    assert summary["A14"].value == "Simple payback (years)"
    assert summary["B14"].value is None
